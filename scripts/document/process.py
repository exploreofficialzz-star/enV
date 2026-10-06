#!/usr/bin/env python3
import sys, json, os, tempfile, shutil, subprocess, zipfile, difflib, re
from pathlib import Path

try:
 import fitz
except Exception: fitz=None
try:
 from pypdf import PdfReader, PdfWriter
except Exception: PdfReader=PdfWriter=None
try:
 from docx import Document
except Exception: Document=None
try:
 from pptx import Presentation
except Exception: Presentation=None
try:
 from openpyxl import load_workbook, Workbook
except Exception: load_workbook=None; Workbook=None
try:
 from PIL import Image
except Exception: Image=None


def fail(msg): raise RuntimeError(msg)
def text_of(path):
    ext=path.suffix.lower()
    if ext=='.pdf':
        if not fitz: fail('PDF text extraction requires PyMuPDF.')
        d=fitz.open(path); return '\n\n'.join(p.get_text() for p in d)
    if ext=='.docx':
        if not Document: fail('DOCX processing requires python-docx.')
        d=Document(path); return '\n'.join(p.text for p in d.paragraphs)
    if ext=='.pptx':
        if not Presentation: fail('PPTX processing requires python-pptx.')
        p=Presentation(path); return '\n\n'.join('\n'.join(sh.text for sh in slide.shapes if hasattr(sh,'text')) for slide in p.slides)
    if ext=='.xlsx':
        if not load_workbook: fail('XLSX processing requires openpyxl.')
        wb=load_workbook(path, data_only=False); rows=[]
        for ws in wb.worksheets:
            rows.append(f'[{ws.title}]')
            for row in ws.iter_rows(values_only=True): rows.append('\t'.join('' if v is None else str(v) for v in row))
        return '\n'.join(rows)
    return path.read_text(errors='replace')

def convert_to_pdf(path, outdir):
    out=Path(outdir); out.mkdir(exist_ok=True)
    r=subprocess.run(['libreoffice','--headless','--convert-to','pdf','--outdir',str(out),str(path)],capture_output=True,text=True,timeout=120)
    if r.returncode!=0: fail(r.stderr[-1000:] or 'LibreOffice conversion failed.')
    pdf=out/(path.stem+'.pdf')
    if not pdf.exists(): fail('LibreOffice did not produce a PDF.')
    return pdf

def pdf_pages(path, indices, output):
    if not PdfReader: fail('PDF page processing requires pypdf.')
    reader=PdfReader(str(path)); writer=PdfWriter()
    for i in indices:
        if 0<=i<len(reader.pages): writer.add_page(reader.pages[i])
    with open(output,'wb') as f: writer.write(f)

def all_images(path, outdir):
    out=Path(outdir); out.mkdir(exist_ok=True); ext=path.suffix.lower(); files=[]
    if ext=='.pdf':
        if not fitz: fail('PDF image extraction requires PyMuPDF.')
        d=fitz.open(path)
        for pi,p in enumerate(d):
            for ii,img in enumerate(p.get_images(full=True)):
                xref=img[0]; data=d.extract_image(xref); fn=out/f'{path.stem}-page{pi+1}-image{ii+1}.{data["ext"]}'; fn.write_bytes(data['image']); files.append(fn)
    elif ext=='.docx':
        with zipfile.ZipFile(path) as z:
            for n in z.namelist():
                if n.startswith('word/media/') and not n.endswith('/'):
                    fn=out/Path(n).name; fn.write_bytes(z.read(n)); files.append(fn)
    elif ext=='.pptx':
        with zipfile.ZipFile(path) as z:
            for n in z.namelist():
                if n.startswith('ppt/media/') and not n.endswith('/'):
                    fn=out/Path(n).name; fn.write_bytes(z.read(n)); files.append(fn)
    elif ext=='.xlsx':
        with zipfile.ZipFile(path) as z:
            for n in z.namelist():
                if n.startswith('xl/media/') and not n.endswith('/'):
                    fn=out/Path(n).name; fn.write_bytes(z.read(n)); files.append(fn)
    else: fail('Image extraction is supported for PDF, DOCX, PPTX and XLSX.')
    return files

def screenshot(path, outdir):
    pdf=path if path.suffix.lower()=='.pdf' else convert_to_pdf(path,outdir)
    out=Path(outdir); subprocess.run(['pdftoppm','-png','-r','144',str(pdf),str(out/'page')],check=True,timeout=120)
    return sorted(out.glob('page-*.png'))

def docx_write_metadata(src,out,title=None,author=None):
    d=Document(src); cp=d.core_properties
    if title is not None: cp.title=title
    if author is not None: cp.author=author
    d.save(out)

def docx_number(src,out):
    d=Document(src)
    for i,p in enumerate(d.paragraphs,1):
        if p.text.strip(): p.text=f'{i}. {p.text}'
    d.save(out)

def docx_watermark(src,out,text):
    d=Document(src)
    d.add_paragraph(f'WATERMARK: {text}')
    d.save(out)

def process(op, inputs, params, work):
    paths=[Path(x) for x in inputs]; first=paths[0]; ext=first.suffix.lower(); out=Path(work)/'output'
    # Cross-format comparisons and text extraction
    if 'comparison' in op:
        if len(paths)<2: fail('Choose two documents to compare.')
        a=text_of(paths[0]).splitlines(); b=text_of(paths[1]).splitlines()
        diff='\n'.join(difflib.unified_diff(a,b,fromfile=paths[0].name,tofile=paths[1].name,linterm=''))
        out.with_suffix('.txt').write_text(diff or 'No textual differences found.'); return out.with_suffix('.txt'),'text/plain'
    if 'text-extractor' in op:
        out.with_suffix('.txt').write_text(text_of(first)); return out.with_suffix('.txt'),'text/plain'
    if 'metadata-tool' in op or 'metadata' in op:
        meta={'name':first.name,'size':first.stat().st_size,'format':ext.lstrip('.')}
        if ext=='.pdf' and fitz:
            d=fitz.open(first); meta.update({'pages':len(d),'metadata':d.metadata})
        elif ext=='.docx' and Document:
            cp=Document(first).core_properties; meta.update({'title':cp.title,'author':cp.author,'subject':cp.subject})
        elif ext=='.pptx' and Presentation: meta['slides']=len(Presentation(first).slides)
        elif ext=='.xlsx' and load_workbook: meta['sheets']=load_workbook(first,read_only=True).sheetnames
        out.with_suffix('.json').write_text(json.dumps(meta,indent=2,ensure_ascii=False)); return out.with_suffix('.json'),'application/json'
    if 'image-text-extractor' in op or op == 'ocr-tool':
        if not Image: fail('OCR requires Pillow.')
        if first.suffix.lower() == '.pdf':
            if not fitz: fail('PDF OCR requires PyMuPDF.')
            doc=fitz.open(first); text=[]
            for i,p in enumerate(doc):
                pix=p.get_pixmap(matrix=fitz.Matrix(1.5,1.5)); img=Image.frombytes('RGB',[pix.width,pix.height],pix.samples); ip=Path(work)/f'ocr-{i+1}.png'; img.save(ip);
                r=subprocess.run(['tesseract',str(ip),'stdout','-l',str(params.get('language','eng'))],capture_output=True,text=True,timeout=120); text.append(r.stdout)
            out.with_suffix('.txt').write_text('\n\n'.join(text)); return out.with_suffix('.txt'),'text/plain'
        r=subprocess.run(['tesseract',str(first),'stdout','-l',str(params.get('language','eng'))],capture_output=True,text=True,timeout=120)
        if r.returncode!=0: fail(r.stderr[-1000:] or 'OCR failed.')
        out.with_suffix('.txt').write_text(r.stdout); return out.with_suffix('.txt'),'text/plain'
    if 'image-splitter' in op and Image:
        img=Image.open(first); cols=max(1,int(params.get('columns',2))); rows=max(1,int(params.get('rows',2))); z=Path(work)/'image-tiles.zip'
        with zipfile.ZipFile(z,'w',zipfile.ZIP_DEFLATED) as zz:
            for y in range(rows):
                for x in range(cols):
                    left=x*img.width//cols; top=y*img.height//rows; right=(x+1)*img.width//cols; bottom=(y+1)*img.height//rows; tile=img.crop((left,top,right,bottom)); tp=Path(work)/f'tile-{y+1}-{x+1}.png'; tile.save(tp); zz.write(tp,tp.name)
        return z,'application/zip'
    if 'image-extractor' in op:
        imgs=all_images(first,Path(work)/'images'); z=Path(work)/'images.zip'
        with zipfile.ZipFile(z,'w',zipfile.ZIP_DEFLATED) as zz:
            for f in imgs: zz.write(f,f.name)
        return z,'application/zip'
    if 'screenshot' in op:
        imgs=screenshot(first,Path(work)/'shots'); z=Path(work)/'screenshots.zip'
        with zipfile.ZipFile(z,'w',zipfile.ZIP_DEFLATED) as zz:
            for f in imgs: zz.write(f,f.name)
        return z,'application/zip'
    if 'page-' in op or op.endswith('splitter') or op.endswith('page-extractor'):
        pdf=first if ext=='.pdf' else convert_to_pdf(first,Path(work)/'converted')
        if not PdfReader: fail('Page operations require pypdf.')
        n=len(PdfReader(str(pdf)).pages); raw=str(params.get('pages') or params.get('range') or '1'); idx=[]
        for token in raw.split(','):
            token=token.strip();
            if not token: continue
            if '-' in token:
                a,b=map(int,token.split('-',1)); idx.extend(range(a-1,b))
            else: idx.append(int(token)-1)
        idx=[i for i in idx if 0<=i<n]
        pdf_pages(pdf,idx,out.with_suffix('.pdf')); return out.with_suffix('.pdf'),'application/pdf'
    if 'compressor' in op:
        z=Path(work)/(first.stem+'-compressed'+first.suffix)
        if ext in ('.docx','.pptx','.xlsx'):
            with zipfile.ZipFile(first) as src, zipfile.ZipFile(z,'w',zipfile.ZIP_DEFLATED,compresslevel=9) as dst:
                for info in src.infolist(): dst.writestr(info,src.read(info.filename))
        elif ext=='.pdf' and PdfReader:
            r=PdfReader(str(first)); w=PdfWriter(); [w.add_page(p) for p in r.pages];
            with open(z,'wb') as f: w.write(f)
        else: shutil.copy2(first,z)
        return z,'application/octet-stream'
    if op=='pdf-to-word':
        if ext!='.pdf': fail('PDF to Word requires a PDF.')
        r=subprocess.run(['libreoffice','--headless','--convert-to','docx','--outdir',str(Path(work)),str(first)],capture_output=True,text=True,timeout=180)
        z=Path(work)/(first.stem+'.docx')
        if r.returncode!=0 or not z.exists(): fail(r.stderr[-1000:] or 'PDF to Word conversion failed.')
        return z,'application/vnd.openxmlformats-officedocument.wordprocessingml.document'
    if ext=='.docx' and Document:
        if 'merger' in op and len(paths)>1:
            outdoc=Document(paths[0])
            for src in paths[1:]:
                d=Document(src)
                for p in d.paragraphs: outdoc.add_paragraph(p.text, style=p.style.name if p.style and p.style.name else None)
                for table in d.tables:
                    t=outdoc.add_table(rows=len(table.rows), cols=len(table.columns))
                    for r,row in enumerate(table.rows):
                        for c,cell in enumerate(row.cells): t.cell(r,c).text=cell.text
            z=Path(work)/(first.stem+'-merged.docx'); outdoc.save(z); return z,'application/vnd.openxmlformats-officedocument.wordprocessingml.document'
        z=Path(work)/(first.stem+'-output.docx')
        if op.endswith('numbering-tool'): docx_number(first,z)
        elif op.endswith('watermark-tool'): docx_watermark(first,z,str(params.get('text','enV')))
        elif op.endswith('print-layout-helper'):
            pdf=convert_to_pdf(first,Path(work)/'print'); return pdf,'application/pdf'
        else: shutil.copy2(first,z)
        return z,'application/vnd.openxmlformats-officedocument.wordprocessingml.document'
    if ext=='.pptx' and Presentation:
        if 'merger' in op and len(paths)>1:
            prs=Presentation(paths[0])
            for src in paths[1:]:
                other=Presentation(src)
                for slide in other.slides:
                    blank=prs.slides.add_slide(prs.slide_layouts[6])
                    for shape in slide.shapes:
                        if not hasattr(shape,'text_frame'): continue
                        box=blank.shapes.add_textbox(shape.left,shape.top,shape.width,shape.height); box.text=shape.text
            z=Path(work)/(first.stem+'-merged.pptx'); prs.save(z); return z,'application/vnd.openxmlformats-officedocument.presentationml.presentation'
        z=Path(work)/(first.stem+'-output.pptx'); shutil.copy2(first,z); return z,'application/vnd.openxmlformats-officedocument.presentationml.presentation'
    if ext=='.xlsx' and load_workbook:
        wb=load_workbook(first)
        if 'merger' in op and len(paths)>1:
            for src in paths[1:]:
                other=load_workbook(src,read_only=True,data_only=False)
                for ws in other.worksheets:
                    target=wb.create_sheet(ws.title[:31])
                    for row in ws.iter_rows():
                        for cell in row:
                            target.cell(cell.row,cell.column).value=cell.value
        z=Path(work)/(first.stem+'-output.xlsx'); wb.save(z); return z,'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    fail(f'Unsupported document operation {op} for {ext}.')

def main():
    if len(sys.argv)<4: fail('Usage: process.py OP OUTPUT_DIR INPUT...')
    op=sys.argv[1]; work=Path(sys.argv[2]); work.mkdir(parents=True,exist_ok=True); meta=json.loads(os.environ.get('ENV_PARAMS','{}')); inputs=sys.argv[3:]
    path,mime=process(op,inputs,meta,work); print(json.dumps({'path':str(path),'mime':mime,'name':path.name}))
if __name__=='__main__':
    try: main()
    except Exception as e: print(json.dumps({'error':str(e)})); sys.exit(2)
